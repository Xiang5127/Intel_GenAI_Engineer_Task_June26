"""Report generator service (Phase 6).

Renders the *normalized* ``report_data`` / ``slide_data`` structures (see
``summarization``) into PDF (ReportLab) and PPTX (python-pptx) files. This layer
is deliberately summarizer-agnostic: it only understands the normalized shape and
does not care whether a rule-based or LLM summarizer produced it.

Normalized shapes consumed:
- report_data: ``{title, subtitle?, metadata{}, sections:[{heading, body, bullets[]}]}``
- slide_data:  ``{title, slides:[{title, bullets[], notes?}]}``
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Optional

# backend/ root (this file is backend/services/report_generator.py).
BACKEND_DIR = Path(__file__).resolve().parents[1]
REPORTS_DIR = BACKEND_DIR / "outputs" / "reports"


class ReportGenerationError(Exception):
    """Raised when a report file cannot be produced."""


def _out_path(prefix: str, ext: str, output_dir: Optional[str]) -> Path:
    base = Path(output_dir) if output_dir else REPORTS_DIR
    base.mkdir(parents=True, exist_ok=True)
    return base / f"{prefix}_{int(time.time())}_{id(object()) % 10000}.{ext}"


def generate_pdf_report(report_data: dict[str, Any], output_dir: Optional[str] = None) -> dict[str, Any]:
    """Render ``report_data`` to a PDF. Returns ``{file_path, file_type, ...}``."""
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    if not isinstance(report_data, dict) or "title" not in report_data:
        raise ReportGenerationError("report_data must be a dict with at least a 'title'")

    out = _out_path("report", "pdf", output_dir)
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("Body", parent=styles["BodyText"], alignment=TA_LEFT, spaceAfter=6)

    story: list[Any] = [Paragraph(str(report_data["title"]), styles["Title"])]
    if report_data.get("subtitle"):
        story.append(Paragraph(str(report_data["subtitle"]), styles["Heading3"]))

    meta = report_data.get("metadata") or {}
    if meta.get("duration_seconds") is not None:
        story.append(Paragraph(f"Duration: {meta['duration_seconds']} s", body_style))
    story.append(Spacer(1, 0.2 * inch))

    for section in report_data.get("sections", []):
        story.append(Paragraph(str(section.get("heading", "")), styles["Heading2"]))
        if section.get("body"):
            story.append(Paragraph(str(section["body"]), body_style))
        bullets = section.get("bullets") or []
        if bullets:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(str(b), body_style)) for b in bullets],
                    bulletType="bullet",
                )
            )
        story.append(Spacer(1, 0.15 * inch))

    try:
        SimpleDocTemplate(str(out), pagesize=LETTER).build(story)
    except Exception as exc:  # noqa: BLE001
        raise ReportGenerationError(f"failed to build PDF: {exc}") from exc

    return {"file_path": str(out), "file_type": "pdf", "sections": len(report_data.get("sections", []))}


def generate_pptx_report(slide_data: dict[str, Any], output_dir: Optional[str] = None) -> dict[str, Any]:
    """Render ``slide_data`` to a PPTX. Returns ``{file_path, file_type, ...}``."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    if not isinstance(slide_data, dict) or "title" not in slide_data:
        raise ReportGenerationError("slide_data must be a dict with at least a 'title'")

    out = _out_path("slides", "pptx", output_dir)
    prs = Presentation()

    # Title slide.
    title_layout = prs.slide_layouts[0]
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = str(slide_data["title"])

    content_layout = prs.slide_layouts[1]
    for slide in slide_data.get("slides", []):
        sl = prs.slides.add_slide(content_layout)
        sl.shapes.title.text = str(slide.get("title", ""))
        body = sl.placeholders[1].text_frame
        body.clear()
        bullets = slide.get("bullets") or []
        if not bullets:
            body.text = ""
        else:
            body.text = str(bullets[0])
            for b in bullets[1:]:
                p = body.add_paragraph()
                p.text = str(b)
                p.level = 1
        if slide.get("notes"):
            sl.notes_slide.notes_text_frame.text = str(slide["notes"])

    try:
        prs.save(str(out))
    except Exception as exc:  # noqa: BLE001
        raise ReportGenerationError(f"failed to build PPTX: {exc}") from exc

    return {"file_path": str(out), "file_type": "pptx", "slides": len(slide_data.get("slides", []))}
